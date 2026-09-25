from datetime import date
from decimal import Decimal
from io import BytesIO
import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from app.db import Base, SessionLocal, engine
from app.models import (
    EntryFieldOption,
    ImportBatch,
    SaleRecord,
    SettlementAfterSaleItem,
    SettlementFeeItem,
    SettlementSummary,
)

from app.schemas import (
    EntryAfterSaleItemCreate,
    EntryCreate,
    EntryFeeItemCreate,
    EntrySaleItemCreate,
)

from app.services.entry_export import build_entry_workbook
from app.services.entry_service import ensure_manual_entry, list_field_options, read_entry, save_entry


@pytest.fixture(autouse=True)


def clean_db():
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    yield


def sale_item(grade="A", head_count="4", sales_quantity=20, unit_price=2.5, **kwargs):
    values = {
        "sale_date": date(2026, 9, 13),
        "variety": kwargs.get("variety", "金枕"),
        "grade": grade,
        "head_count": head_count,
        "spec_kg": kwargs.get("spec_kg", "10"),
        "sales_quantity": Decimal(str(sales_quantity)),
        "unit_price": Decimal(str(unit_price)),
        "remark": kwargs.get("remark"),
    }
    return EntrySaleItemCreate(**values)


def after_item(content="售后", summary="摘要", amount=10):
    return EntryAfterSaleItemCreate(
        content=content,
        summary=summary,
        amount=Decimal(str(amount)),
    )


def fee_item(name="代卖佣金", amount=5, is_custom=False):
    return EntryFeeItemCreate(
        name=name,
        amount=Decimal(str(amount)),
        is_custom=is_custom,
    )


def entry_payload(overwrite=False, sales=None, after_sales=None, fees=None, **kwargs):
    values = {
        "merchant_no": "637",
        "order_no": "宝贝-001",
        "container_no": "C001",
        "vehicle_no": "桂A0001",
        "country": "越南",
        "market": "南宁海吉星",
        "arrival_date": date(2026, 9, 10),
        "arrival_quantity": 20,
        "sales": sales if sales is not None else [sale_item()],
        "after_sales": after_sales if after_sales is not None else [],
        "fees": fees if fees is not None else [],
        "overwrite": overwrite,
    }
    values.update(kwargs)
    return EntryCreate(**values)


SALES_HEADERS = ("销售日期", "品种", "数量(件)", "单价(元)", "金额(元)")


def _sheet_row_containing(ws, text):
    """按任意单元格文字定位行号，避免把导出布局的行号写死在断言里。"""

    for row in range(1, ws.max_row + 1):
        for cell in ws[row]:
            if isinstance(cell.value, str) and text in cell.value:
                return row
    raise AssertionError(f"导出结果中没有包含「{text}」的行")


def _sales_header(ws):
    """定位销售明细表头，返回行号与「表头文字 → 列号」映射。"""

    for row in range(1, ws.max_row + 1):
        columns = {str(cell.value): cell.column for cell in ws[row] if cell.value}
        if set(SALES_HEADERS) <= set(columns):
            return row, columns
    raise AssertionError("导出结果中没有找到销售明细表头")


def _number(value):
    """兼容数值单元格与千分位文本单元格。"""

    if value is None:
        return None
    if isinstance(value, int | float | Decimal):
        return float(value)
    return float(str(value).replace(",", ""))


def test_save_and_read_entry_persists_sales_after_sales_fees_and_summary():
    db = SessionLocal()
    payload = entry_payload(
        sales=[
            sale_item("A", "4", 20, 2.5),
            sale_item("B", "6/8", 30, 3, spec_kg="10"),
        ],
        after_sales=[after_item("坏果", "扣款", 10), after_item("补货", "", 20)],
        fees=[
            fee_item("代卖佣金", 5),
            fee_item("运费", 10),
            fee_item("其他", 3, is_custom=True),
        ],
    )
    result = save_entry(db, payload)
    assert result.status == "created"
    assert result.merchant_no == "637"
    batch = db.query(ImportBatch).one()
    assert batch.source_type == "manual"
    assert batch.market == "南宁海吉星"
    assert batch.arrival_date == date(2026, 9, 10)
    assert batch.arrival_quantity == 20
    records = db.query(SaleRecord).order_by(SaleRecord.id).all()
    assert [record.grade.value for record in records] == ["A", "B"]
    assert [record.piece_count for record in records] == ["4", "6/8"]
    assert [record.piece_count_min for record in records] == [Decimal("4"), Decimal("6")]
    assert [record.piece_count_max for record in records] == [Decimal("4"), Decimal("8")]
    assert [record.spec_kg for record in records] == ["10", "10"]
    assert [record.spec_kg_max for record in records] == [Decimal("10"), Decimal("10")]
    assert [record.amount for record in records] == [Decimal("50.00"), Decimal("90.00")]
    assert db.query(SettlementAfterSaleItem).count() == 2
    assert db.query(SettlementFeeItem).count() == 3
    summary = db.query(SettlementSummary).one()
    assert summary.sales_amount == Decimal("140.00")
    assert summary.after_sale_amount == Decimal("30.00")
    assert summary.goods_amount == Decimal("110.00")
    assert summary.fee_amount == Decimal("18.00")
    assert summary.payable_amount == Decimal("92.00")
    entry = read_entry(db, "637")
    assert entry is not None
    assert entry["source_type"] == "manual"
    assert len(entry["sales"]) == 2
    assert entry["after_sales"][0]["amount"] == Decimal("10.00")
    assert entry["fees"][-1]["is_custom"] is True
    db.close()


def test_save_entry_allows_blank_variety_spec_and_zero_price():
    db = SessionLocal()
    payload = entry_payload(
        sales=[sale_item("", "", 5, 0, spec_kg="", remark="只填备注和数量", variety="")],
    )
    result = save_entry(db, payload)
    assert result.status == "created"
    record = db.query(SaleRecord).one()
    assert record.grade.value == "OTHER"
    assert record.variety is None
    assert record.piece_count is None
    assert record.spec_kg is None
    assert record.unit_price == Decimal("0.00")
    assert record.amount == Decimal("0.00")
    entry = read_entry(db, "637")
    assert entry is not None
    assert entry["sales"][0]["variety"] == ""
    assert entry["sales"][0]["grade"] == "OTHER"
    assert entry["sales"][0]["head_count"] == ""
    assert entry["sales"][0]["spec_kg"] == ""
    db.close()


def test_sale_item_still_rejects_non_empty_invalid_spec():
    with pytest.raises(ValidationError):
        EntrySaleItemCreate(
            sale_date=date(2026, 9, 13),
            variety="A",
            head_count="abc",
            spec_kg="",
            sales_quantity=Decimal("5"),
            unit_price=Decimal("0"),
        )


def test_sale_item_rejects_kg_range():
    with pytest.raises(ValidationError):
        EntrySaleItemCreate(
            sale_date=date(2026, 9, 13),
            variety="A",
            head_count="4",
            spec_kg="9/10",
            sales_quantity=Decimal("5"),
            unit_price=Decimal("0"),
        )


def test_save_without_overwrite_returns_conflict_and_overwrite_replaces():
    db = SessionLocal()
    first = entry_payload()
    save_entry(db, first)
    conflict = save_entry(db, entry_payload(order_no="新单号"))
    assert conflict.status == "conflict"
    assert conflict.existing_order_no == "宝贝-001"
    replaced = save_entry(
        db,
        entry_payload(
            overwrite=True,
            order_no="新单号",
            sales=[sale_item("C", "2", 5, 10), sale_item("D", "3", 7, 8)],
        ),
    )
    assert replaced.status == "saved"
    assert db.query(ImportBatch).count() == 1
    assert db.query(SaleRecord).count() == 2
    assert db.query(SettlementSummary).one().sales_amount == Decimal("106.00")
    assert db.query(ImportBatch).one().order_no == "新单号"
    manual = ensure_manual_entry(db, "637")
    assert manual.source_type == "manual"
    with pytest.raises(ValueError):
        ensure_manual_entry(db, "missing")
    db.close()


def test_manual_save_refuses_to_overwrite_import_batch():
    db = SessionLocal()
    db.add(
        ImportBatch(
            merchant_no="637",
            merchant_no_normalized="637",
            source_type="import",
            status="success",
            success_count=1,
            warning_count=0,
            failure_count=0,
        )
    )
    db.commit()
    result = save_entry(db, entry_payload(overwrite=True))
    assert result.status == "conflict"
    assert result.existing_source_type == "import"
    assert "导入数据占用" in (result.conflict_reason or "")
    assert db.query(ImportBatch).filter_by(merchant_no="637").one().source_type == "import"
    db.close()


def test_manual_save_rejects_different_raw_merchant_with_same_normalized_value():
    db = SessionLocal()
    db.add(
        ImportBatch(
            merchant_no="单637",
            merchant_no_normalized="637",
            source_type="manual",
            status="success",
            success_count=1,
            warning_count=0,
            failure_count=0,
        )
    )
    db.commit()
    result = save_entry(db, entry_payload(merchant_no="637", overwrite=True))
    assert result.status == "conflict"
    assert "归一化后与已有商号 单637 相同" in (result.conflict_reason or "")
    assert db.query(ImportBatch).filter_by(merchant_no="637").count() == 0
    db.close()


def test_field_options_only_return_active_rows_in_order():
    db = SessionLocal()
    db.add_all([
        EntryFieldOption(field_key="market", value="市场二", sort_order=2, is_active=True),
        EntryFieldOption(field_key="market", value="市场一", sort_order=1, is_active=True),
        EntryFieldOption(field_key="market", value="停用市场", sort_order=0, is_active=False),
        EntryFieldOption(field_key="variety", value="C", sort_order=2, is_active=True),
    ])
    db.commit()
    assert [item["value"] for item in list_field_options(db, "market")] == ["市场一", "市场二"]
    assert [item["value"] for item in list_field_options(db, "variety")] == ["C"]
    db.close()


def test_export_handles_dynamic_rows_and_writes_only_values():
    db = SessionLocal()
    payload = entry_payload(
        sales=[
            sale_item("A", "4", 20, 2.5),
            sale_item("B", "6/8", 30, 3, spec_kg="10"),
            sale_item("C", "2", 10, 4),
        ],
        after_sales=[
            after_item("内容1", "摘要1", 1),
            after_item("内容2", "摘要2", 2),
            after_item("内容3", "摘要3", 3),
            after_item("内容4", "摘要4", 4),
            after_item("内容5", "摘要5", 5),
        ],
        fees=[
            fee_item("代卖佣金", 5),
            fee_item("运费", 10),
            fee_item("车位费", 11),
            fee_item("入场费", 12),
            fee_item("搬运费", 13),
            fee_item("打冷费", 14),
            fee_item("临时人工", 9, is_custom=True),
            fee_item("临时车费", 7, is_custom=True),
        ],
    )
    save_entry(db, payload)
    content = build_entry_workbook(db, "637")
    wb = load_workbook(BytesIO(content))
    ws = wb["结算单"]
    assert ws["A1"].value == "结 算 单"
    info = ws["A3"].value
    assert "商号：637" in info
    assert "市场：南宁海吉星" in info
    assert "到达日期：2026-09-10" in info
    assert "来货数量：20" in info
    sales_header, columns = _sales_header(ws)
    assert {"品种", "等级", "规格(头数)", "规格(KG)", "备注", "数量(件)", "单价(元)", "金额(元)"} <= set(columns)
    # 三条明细按录入顺序落行，金额一律按「数量 × 单价」重算。
    assert [
        _number(ws.cell(sales_header + offset, columns["数量(件)"]).value) for offset in (1, 2, 3)
    ] == [20.0, 30.0, 10.0]
    assert [
        _number(ws.cell(sales_header + offset, columns["金额(元)"]).value) for offset in (1, 2, 3)
    ] == [50.0, 90.0, 40.0]
    total_row = _sheet_row_containing(ws, "总件数")
    assert _number(ws.cell(total_row, 7).value) == 60
    assert _number(ws.cell(total_row, 9).value) == 180
    after_total_row = _sheet_row_containing(ws, "售后合计")
    assert _number(ws.cell(after_total_row, 7).value) == 15
    goods_row = _sheet_row_containing(ws, "货款合计")
    assert _number(ws.cell(goods_row, 8).value) == 165
    fee_total_row = _sheet_row_containing(ws, "费用合计")
    # 自定义费用只能出现一次，否则费用明细行与小计自相矛盾。
    fee_names = [
        ws.cell(row, 1).value for row in range(1, fee_total_row) if isinstance(ws.cell(row, 1).value, str)
    ]
    assert fee_names.count("临时人工") == 1
    assert fee_names.count("临时车费") == 1
    assert _number(ws.cell(fee_total_row, 7).value) == 81
    payable_row = _sheet_row_containing(ws, "应付贵方总金额")
    assert _number(ws.cell(payable_row, 7).value) == 84
    formula_cells = [
        cell.coordinate
        for row in ws.iter_rows()
        for cell in row
        if cell.data_type == "f"
    ]
    assert formula_cells == []
    db.close()


def test_export_dynamic_rows_keep_consistent_style_and_height():
    """同一区块内动态生成的行必须样式、行高一致，避免打印版式走样。"""

    db = SessionLocal()
    payload = entry_payload(
        sales=[
            sale_item("A", "4", 20, 2.5),
            sale_item("B", "6/8", 30, 3, spec_kg="10"),
            sale_item("C", "2", 10, 4),
            sale_item("D", "3", 15, 5),
        ],
        after_sales=[after_item(f"内容{i}", f"摘要{i}", i) for i in range(1, 7)],
        fees=[
            fee_item("代卖佣金", 5),
            fee_item("运费", 10),
            fee_item("车位费", 11),
            fee_item("入场费", 12),
            fee_item("搬运费", 13),
            fee_item("打冷费", 14),
            fee_item("临时人工", 9, is_custom=True),
        ],
    )
    save_entry(db, payload)
    ws = load_workbook(BytesIO(build_entry_workbook(db, "637")))["结算单"]

    def style_of(cell):
        return (
            cell.font.sz,
            cell.number_format,
            cell.alignment.horizontal,
            cell.border.bottom.style,
        )

    def assert_rows_consistent(first_row, other_rows):
        for row in other_rows:
            for column in (1, 2, 7, 8):
                assert style_of(ws.cell(row, column)) == style_of(ws.cell(first_row, column))
            assert ws.row_dimensions[row].height == ws.row_dimensions[first_row].height

    sales_header, _ = _sales_header(ws)
    assert_rows_consistent(sales_header + 1, range(sales_header + 2, sales_header + 5))

    after_header = _sheet_row_containing(ws, "序号")
    assert_rows_consistent(after_header + 1, range(after_header + 2, after_header + 7))

    fee_header = _sheet_row_containing(ws, "费用项目")
    fee_total_row = _sheet_row_containing(ws, "费用合计")
    assert_rows_consistent(fee_header + 1, range(fee_header + 2, fee_total_row))

    # 导出只写值不写公式，避免客户打开后触发重算或引用失效。
    assert [
        cell.coordinate for row in ws.iter_rows() for cell in row if cell.data_type == "f"
    ] == []
    db.close()
