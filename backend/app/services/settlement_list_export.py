"""数据明细「结算单列表」导出：结算单汇总 + 销售 / 售后 / 支出费用分类明细。"""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from ..models import (
    ImportBatch,
    SaleRecord,
    SettlementAfterSaleItem,
    SettlementFeeItem,
    SettlementSummary,
)
from .analytics_core import GRADES
from .field_conversion import grade_mapping_note
from .merchant_no_naming import merchant_no_display
from .order_no_naming import order_no_display
from .settlement_list_service import list_settlements


SHEET_ROWS = "结算单列表"
SHEET_SALES = "销售明细"
SHEET_AFTER_SALE = "售后明细"
SHEET_FEES = "支出费用明细"
SHEET_NOTES = "说明"
EMPTY_NOTE = "当前筛选范围没有结算单"
EMPTY_SALES_NOTE = "当前筛选范围没有销售明细"
EMPTY_AFTER_SALE_NOTE = "当前筛选范围没有售后明细"
EMPTY_FEES_NOTE = "当前筛选范围没有支出费用明细"
NUMBER_FORMAT = "#,##0.00"
INTEGER_FORMAT = "#,##0"
DATE_FORMAT = "yyyy-mm-dd"
HEADER_FILL = PatternFill("solid", fgColor="E7F1EB")
BLANK = "—"
SUMMARY_AFTER_SALE_CONTENT = "结算摘要售后合计"
SUMMARY_FEE_SOURCE = "结算摘要"
MANUAL_FEE_SOURCE = "录单录入"
MANUAL_AFTER_SALE_SOURCE = "录单录入"
GRADE_LABELS = {
    "A": "A果",
    "B": "B果",
    "AB": "AB果",
    "C": "C果",
    "D": "D果",
    "E": "E果",
    "F": "F果",
    "OTHER": "其他",
}


def _cell(value):
    """Excel 会把 =+-@ 开头的文本当公式执行，这里沿用 CSV 导出的转义口径。"""

    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{value}"
    return value


def _decimal(value: str) -> Decimal | None:
    text = (value or "").strip().replace(",", "")
    if not text:
        return None
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def _fee_name(label: str) -> str:
    """费用明细文本里夹带了原始金额单元格，这里只保留费用名称。"""

    parts = label.split()
    while len(parts) > 1 and _decimal(parts[-1]) is not None:
        parts.pop()
    return " ".join(parts).strip() or label.strip()


def parse_fee_detail(text: str | None) -> list[tuple[str, Decimal]]:
    """把结算摘要的『费用名 原值: 金额；…』文本拆成 (费用项目, 金额) 列表。"""

    if not text:
        return []
    rows: list[tuple[str, Decimal]] = []
    for segment in text.split("；"):
        label, separator, amount_text = segment.rpartition(":")
        if not separator:
            continue
        name = _fee_name(label)
        amount = _decimal(amount_text)
        if name and amount is not None:
            rows.append((name, amount))
    return rows


def _context(batch: ImportBatch) -> tuple:
    """明细行前缀：商号 / 单号 / 柜号，与列表页显示口径一致。"""

    return (
        merchant_no_display(batch.merchant_no, batch.merchant_no_normalized) or BLANK,
        order_no_display(batch.order_no, batch.order_no_normalized) or BLANK,
        batch.container_no or BLANK,
    )


def _columns(grade_codes: list[str]) -> list[tuple[str, str, int]]:
    """返回 (表头, 数字格式, 列宽)；数字格式为空串表示文本列。"""

    columns = [
        ("商号", "", 16),
        ("品牌", "", 14),
        ("单号", "", 16),
        ("柜号", "", 16),
        ("销售日期起", DATE_FORMAT, 14),
        ("销售日期止", DATE_FORMAT, 14),
        ("总件数", NUMBER_FORMAT, 12),
    ]
    columns += [
        (f"{GRADE_LABELS.get(code, code)}件数", NUMBER_FORMAT, 12) for code in grade_codes
    ]
    columns += [
        ("销售金额", NUMBER_FORMAT, 14),
        ("每件均价", NUMBER_FORMAT, 16),
        ("售后合计", NUMBER_FORMAT, 14),
        ("费用合计", NUMBER_FORMAT, 14),
        ("应付贵方总金额(RMB)", NUMBER_FORMAT, 20),
    ]
    return columns


def _summary_amount(summary: SettlementSummary | None, field: str):
    value = getattr(summary, field, None) if summary is not None else None
    if field == "after_sale_amount" and value is not None:
        value = abs(value)
    return "" if value is None else value


def _row(item: dict, grade_codes: list[str], summary: SettlementSummary | None) -> list:
    quantities = item["grade_quantities"]
    values = [
        item["merchant_no_normalized"] or item["merchant_no"] or BLANK,
        item.get("brand") or BLANK,
        item["order_no_normalized"] or item["order_no"] or BLANK,
        item["container_no"] or BLANK,
        item["sale_date_start"],
        item["sale_date_end"],
        item["total_quantity"],
    ]
    values += [quantities.get(code) or 0 for code in grade_codes]
    values += [
        item["sales_amount"],
        item["average_price"] if item["average_price"] is not None else "",
        _summary_amount(summary, "after_sale_amount"),
        _summary_amount(summary, "fee_amount"),
        _summary_amount(summary, "payable_amount"),
    ]
    return [_cell(value) for value in values]


def _visible_grades(items: list[dict]) -> list[str]:
    """与页面口径一致：只导出当前范围内真正有数量的等级列。"""

    return [
        grade.value
        for grade in GRADES
        if any(float(item["grade_quantities"].get(grade.value) or 0) > 0 for item in items)
    ]


def _sales_columns() -> list[tuple[str, str, int]]:
    return [
        ("商号", "", 16),
        ("单号", "", 16),
        ("柜号", "", 16),
        ("销售日期", DATE_FORMAT, 14),
        ("品种", "", 16),
        ("等级原文", "", 18),
        ("标准等级", "", 10),
        ("规格原文", "", 20),
        ("规格（头数）", INTEGER_FORMAT, 14),
        ("规格（KG）", NUMBER_FORMAT, 14),
        ("销售数量", NUMBER_FORMAT, 14),
        ("单价", NUMBER_FORMAT, 12),
        ("金额", NUMBER_FORMAT, 14),
        ("备注", "", 28),
    ]


def _after_sale_columns() -> list[tuple[str, str, int]]:
    return [
        ("商号", "", 16),
        ("单号", "", 16),
        ("柜号", "", 16),
        ("内容", "", 26),
        ("摘要", "", 30),
        ("金额", NUMBER_FORMAT, 14),
        ("来源", "", 14),
    ]


def _fee_columns() -> list[tuple[str, str, int]]:
    return [
        ("商号", "", 16),
        ("单号", "", 16),
        ("柜号", "", 16),
        ("费用项目", "", 22),
        ("金额", NUMBER_FORMAT, 14),
        ("来源", "", 14),
    ]


def _sales_rows(batch: ImportBatch, records: list[SaleRecord]) -> list[list]:
    prefix = _context(batch)
    return [
        [
            _cell(prefix[0]),
            _cell(prefix[1]),
            _cell(prefix[2]),
            record.sale_date,
            _cell(record.variety or BLANK),
            _cell(record.grade_raw or BLANK),
            record.grade.value if record.grade is not None else BLANK,
            _cell(record.spec_raw or BLANK),
            record.piece_count if record.piece_count is not None else "",
            record.spec_kg if record.spec_kg is not None else "",
            record.quantity,
            record.unit_price,
            record.amount,
            _cell(record.remark or ""),
        ]
        for record in records
    ]


def _after_sale_rows(
    batch: ImportBatch,
    items: list[SettlementAfterSaleItem],
    summary: SettlementSummary | None,
) -> list[list]:
    prefix = _context(batch)
    if items:
        return [
            [
                _cell(prefix[0]),
                _cell(prefix[1]),
                _cell(prefix[2]),
                _cell(item.content),
                _cell(item.summary),
                abs(item.amount) if item.amount is not None else BLANK,
                MANUAL_AFTER_SALE_SOURCE,
            ]
            for item in items
        ]
    amount = getattr(summary, "after_sale_amount", None) if summary is not None else None
    if amount is None:
        return []
    amount = abs(amount)
    return [
        [
            _cell(prefix[0]),
            _cell(prefix[1]),
            _cell(prefix[2]),
            SUMMARY_AFTER_SALE_CONTENT,
            "导入结算摘要原值",
            amount,
            SUMMARY_FEE_SOURCE,
        ]
    ]


def _fee_rows(
    batch: ImportBatch,
    items: list[SettlementFeeItem],
    summary: SettlementSummary | None,
) -> list[list]:
    prefix = _context(batch)
    if items:
        return [
            [
                _cell(prefix[0]),
                _cell(prefix[1]),
                _cell(prefix[2]),
                _cell(item.name),
                item.amount,
                "录单自定义" if item.is_custom else MANUAL_FEE_SOURCE,
            ]
            for item in items
        ]
    if summary is None:
        return []
    parsed = parse_fee_detail(summary.fee_detail)
    if parsed:
        return [
            [_cell(prefix[0]), _cell(prefix[1]), _cell(prefix[2]), _cell(name), amount, SUMMARY_FEE_SOURCE]
            for name, amount in parsed
        ]
    if summary.fee_amount is None:
        return []
    return [
        [
            _cell(prefix[0]),
            _cell(prefix[1]),
            _cell(prefix[2]),
            "结算摘要费用合计",
            summary.fee_amount,
            SUMMARY_FEE_SOURCE,
        ]
    ]


def _write_table_sheet(
    sheet,
    columns: list[tuple[str, str, int]],
    rows: list[list],
    empty_note: str,
) -> None:
    for index, (title, _, width) in enumerate(columns, start=1):
        header = sheet.cell(1, index, title)
        header.font = Font(bold=True)
        header.fill = HEADER_FILL
        header.alignment = Alignment(horizontal="center")
        sheet.column_dimensions[get_column_letter(index)].width = width
    if not rows:
        sheet.cell(2, 1, empty_note)
    for offset, row in enumerate(rows, start=2):
        for index, value in enumerate(row, start=1):
            cell = sheet.cell(offset, index, value)
            number_format = columns[index - 1][1]
            if number_format:
                cell.number_format = number_format
    sheet.freeze_panes = "A2"
    if rows:
        last_column = get_column_letter(len(columns))
        sheet.auto_filter.ref = f"A1:{last_column}{len(rows) + 1}"


def _batch_index(db: Session, items: list[dict]) -> dict[str, ImportBatch]:
    """按商号取回结算单批次，明细与汇总表复用同一份批次。"""

    merchants = [item["merchant_no"] for item in items if item["merchant_no"]]
    if not merchants:
        return {}
    batches = db.query(ImportBatch).filter(ImportBatch.merchant_no.in_(merchants)).all()
    return {batch.merchant_no: batch for batch in batches}


def _detail_source(db: Session, batches: list[ImportBatch], date_range: dict | None) -> dict:
    """收集明细所需的销售 / 售后 / 费用行与结算摘要，统一按批次分组。"""

    empty = {"records": {}, "after_sales": {}, "fees": {}, "summaries": {}}
    if not batches or date_range is None:
        return empty
    batch_ids = [batch.id for batch in batches]
    records = (
        db.query(SaleRecord)
        .filter(
            SaleRecord.import_batch_id.in_(batch_ids),
            SaleRecord.sale_date >= date_range["start_date"],
            SaleRecord.sale_date <= date_range["end_date"],
        )
        .order_by(SaleRecord.import_batch_id, SaleRecord.sale_date, SaleRecord.id)
        .all()
    )
    after_sales = (
        db.query(SettlementAfterSaleItem)
        .filter(SettlementAfterSaleItem.import_batch_id.in_(batch_ids))
        .order_by(SettlementAfterSaleItem.import_batch_id, SettlementAfterSaleItem.sort_order)
        .all()
    )
    fees = (
        db.query(SettlementFeeItem)
        .filter(SettlementFeeItem.import_batch_id.in_(batch_ids))
        .order_by(SettlementFeeItem.import_batch_id, SettlementFeeItem.sort_order)
        .all()
    )
    summaries = {
        summary.import_batch_id: summary
        for summary in db.query(SettlementSummary)
        .filter(SettlementSummary.import_batch_id.in_(batch_ids))
        .all()
    }
    return {
        "records": _by_batch(records),
        "after_sales": _by_batch(after_sales),
        "fees": _by_batch(fees),
        "summaries": summaries,
    }


def _by_batch(rows: list) -> dict[int, list]:
    grouped: dict[int, list] = defaultdict(list)
    for row in rows:
        grouped[row.import_batch_id].append(row)
    return grouped


def _write_notes_sheet(
    sheet,
    data: dict,
    batches: list[ImportBatch],
    counts: dict[str, int],
    merchant_no: str | None,
    brand: str | None,
    grade_mapping: str,
) -> None:
    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 64
    date_range = data["date_range"]
    rows = [
        ("导出内容", "结算单列表（汇总）+ 销售明细 / 售后明细 / 支出费用明细（分类明细）"),
        ("销售日期起", date_range["start_date"].isoformat() if date_range else "暂无销售数据"),
        ("销售日期止", date_range["end_date"].isoformat() if date_range else "暂无销售数据"),
        (
            "是否默认范围",
            "是（最新销售日期往前一个月）" if date_range and date_range["is_default"] else "否",
        ),
        ("筛选商号", merchant_no or "全部结算单"),
        ("筛选品牌", brand or "全部品牌"),
        ("导出行数", len(batches)),
        ("销售明细行数", counts["sales"]),
        ("售后明细行数", counts["after_sales"]),
        ("支出费用行数", counts["fees"]),
        ("生成时间(UTC)", datetime.now(timezone.utc).isoformat()),
        ("等级映射", f"{grade_mapping}；明细等级原文保留在 grade_raw"),
        ("件数口径", "总件数与各等级件数为销售数量合计（千克）"),
        ("金额口径", "销售金额为明细金额合计；每件均价 = 销售金额 ÷ 总件数"),
        ("售后合计 / 费用合计 / 应付贵方总金额", "取自结算摘要原表；手工录单件按录单页口径计算"),
        ("销售明细口径", "每条销售记录一行；规格（头数）/ 规格（KG）仅手工录单有值，导入件看『规格原文』"),
        ("售后明细口径", "手工录单售后行；导入件没有明细时回填结算摘要的售后合计原值"),
        ("支出费用口径", "手工录单费用行（含自定义项）；导入件按结算摘要的费用明细文本拆分为『费用项目 + 金额』"),
        ("导出口径", "导出当前筛选范围的全部结算单，与列表页分页无关"),
    ]
    for offset, (label, value) in enumerate(rows, start=1):
        sheet.cell(offset, 1, _cell(label)).font = Font(bold=True)
        sheet.cell(offset, 2, _cell(value))


def build_settlements_workbook(
    db: Session,
    *,
    start_date: date | None = None,
    end_date: date | None = None,
    merchant_no: str | None = None,
    brand: str | None = None,
) -> bytes:
    """导出「结算单列表」当前筛选结果（汇总 + 分类明细），返回 xlsx 字节。"""

    data = list_settlements(
        db,
        start_date=start_date,
        end_date=end_date,
        merchant_no=merchant_no,
        brand=brand,
    )
    items = data["settlements"]
    grade_codes = _visible_grades(items)
    batch_index = _batch_index(db, items)
    batches = [
        batch_index[item["merchant_no"]]
        for item in items
        if item["merchant_no"] in batch_index
    ]
    source = _detail_source(db, batches, data["date_range"])

    summary_rows = []
    sales_rows: list[list] = []
    after_sale_rows: list[list] = []
    fee_rows: list[list] = []
    for item in items:
        batch = batch_index.get(item["merchant_no"])
        summary = source["summaries"].get(batch.id) if batch is not None else None
        summary_rows.append(_row(item, grade_codes, summary))
        if batch is None:
            continue
        sales_rows += _sales_rows(batch, source["records"].get(batch.id, []))
        after_sale_rows += _after_sale_rows(
            batch, source["after_sales"].get(batch.id, []), summary
        )
        fee_rows += _fee_rows(batch, source["fees"].get(batch.id, []), summary)

    workbook = Workbook()
    _write_table_sheet(workbook.active, _columns(grade_codes), summary_rows, EMPTY_NOTE)
    workbook.active.title = SHEET_ROWS
    _write_table_sheet(
        workbook.create_sheet(SHEET_SALES), _sales_columns(), sales_rows, EMPTY_SALES_NOTE
    )
    _write_table_sheet(
        workbook.create_sheet(SHEET_AFTER_SALE),
        _after_sale_columns(),
        after_sale_rows,
        EMPTY_AFTER_SALE_NOTE,
    )
    _write_table_sheet(
        workbook.create_sheet(SHEET_FEES), _fee_columns(), fee_rows, EMPTY_FEES_NOTE
    )
    _write_notes_sheet(
        workbook.create_sheet(SHEET_NOTES),
        data,
        batches,
        {"sales": len(sales_rows), "after_sales": len(after_sale_rows), "fees": len(fee_rows)},
        merchant_no,
        brand,
        grade_mapping_note(db),
    )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def settlements_export_filename(
    start_date: date | None,
    end_date: date | None,
    merchant_no: str | None,
    brand: str | None,
) -> str:
    parts = ["结算单列表"]
    if merchant_no:
        parts.append(merchant_no)
    if brand:
        parts.append(brand)
    if start_date and end_date:
        parts.append(f"{start_date:%Y%m%d}-{end_date:%Y%m%d}")
    return "-".join(parts) + ".xlsx"


__all__ = [
    "build_settlements_workbook",
    "parse_fee_detail",
    "settlements_export_filename",
]
